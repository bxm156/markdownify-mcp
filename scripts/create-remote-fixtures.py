"""Generate real Office documents for parser smoke tests, without external data.

Usage: python scripts/create-remote-fixtures.py DESTINATION_DIR

Writes test.docx, test.xlsx and test.pptx into DESTINATION_DIR (created if
needed). Each holds "Test <EXT> content" and "Office parser verification",
which remote-converter-smoke.ts expects after conversion. Requires openpyxl
and python-pptx. A missing argument prints the usage line and exits 2.
"""
import sys

if len(sys.argv) != 2:
    print("usage: python scripts/create-remote-fixtures.py DESTINATION_DIR", file=sys.stderr)
    sys.exit(2)

from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from openpyxl import Workbook
from pptx import Presentation

destination = Path(sys.argv[1])
destination.mkdir(parents=True, exist_ok=True)
# DOCX built by hand (no python-docx dependency): the smallest valid package,
# i.e. content types, the package relationship and the main document part.
with ZipFile(destination / "test.docx", "w", ZIP_DEFLATED) as document:
    document.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    document.writestr("_rels/.rels", '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    document.writestr("word/document.xml", '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Test DOCX content</w:t></w:r></w:p><w:p><w:r><w:t>Office parser verification</w:t></w:r></w:p><w:sectPr/></w:body></w:document>')
workbook = Workbook()
workbook.active.append(["Test XLSX content", "Office parser verification"])
workbook.active.append(["Quantity", 42])
workbook.save(destination / "test.xlsx")
presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[1])
slide.shapes.title.text = "Test PPTX content"
slide.placeholders[1].text = "Office parser verification"
presentation.save(destination / "test.pptx")
print("Generated real OOXML DOCX, XLSX and PPTX fixtures")
